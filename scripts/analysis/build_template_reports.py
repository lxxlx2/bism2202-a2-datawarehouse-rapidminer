"""Fill the immutable teacher template; FINAL mode requires genuine native images."""
from pathlib import Path
from copy import deepcopy
from io import BytesIO
import argparse,hashlib,json,zipfile
from lxml import etree as E
from docx import Document
from docx.shared import Inches, Pt
ROOT=Path(__file__).resolve().parents[2]
REFERENCE=ROOT/'requirements/original/Assginmen2_Case_BISM2202_2026s2_Business_Report.docx'
W='http://schemas.openxmlformats.org/wordprocessingml/2006/main'
WP='http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing'
A='http://schemas.openxmlformats.org/drawingml/2006/main'
NS={'w':W,'wp':WP,'a':A}
def q(n):return '{'+W+'}'+n
def prose(doc,text):
 for part in text.split('\n\n'):
  p=doc.add_paragraph(part,style='Normal');p.paragraph_format.space_after=None
 return p

def sqlbox(doc,source,sql,slot):
 # Editable SQL slots use a native Word cell so both Word and LibreOffice
 # paginate the complete answer without floating-shape clipping.
 lines=[l.rstrip() for l in sql.splitlines() if not l.startswith('--') and l.strip()]
 table=doc.add_table(rows=1,cols=1);table.autofit=False
 table.columns[0].width=Inches(6.25)
 pr=table._tbl.tblPr
 borders=E.SubElement(pr,q('tblBorders'))
 for edge in ['top','left','bottom','right']:
  E.SubElement(borders,q(edge),{q('val'):'single',q('sz'):'10',q('color'):'7030A0'})
 cell=table.cell(0,0)
 for i,line in enumerate(lines):
  p=cell.paragraphs[0] if i==0 else cell.add_paragraph()
  p.style='Normal';p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0)
  p.paragraph_format.keep_with_next=i<len(lines)-1
  r=p.add_run(line);r.font.name='Consolas';r.font.size=Pt(10.5)
 doc.add_paragraph().paragraph_format.space_after=Pt(0)

IMAGE_SLOTS={61:['schema/STUDENT_{s}_ID_dm.png'],88:['gui/ssis_dimcustomer.png','gui/ssis_dimproduct.png','gui/ssis_dimseller.png','gui/ssis_dimdate.png','gui/ssis_geography_customer.png','gui/ssis_geography_seller.png'],101:['gui/ssis_fact.png'],138:['gui/sql_Q4_1.png'],170:['gui/sql_Q4_2.png'],212:['gui/sql_Q4_3.png'],263:['gui/model1_process.png','gui/model1_cv.png'],267:['gui/model2_process.png','gui/model2_cv.png'],284:['gui/model1_result.png'],292:['gui/model2_result.png'],303:['gui/model1_performance.png','gui/model2_performance.png']}
LABELS={'ssis_dimcustomer':'Customer dimension native SSIS Data Flow','ssis_dimproduct':'Product dimension native SSIS Data Flow','ssis_dimseller':'Seller dimension native SSIS Data Flow','ssis_dimdate':'Date dimension native SSIS Data Flow','ssis_geography_customer':'Customer destination Geography native SSIS Data Flow','ssis_geography_seller':'Supplier origin Geography native SSIS Data Flow','ssis_fact':'FactSales native SSIS Data Flow','sql_Q4_1':'Executed customer age revenue query','sql_Q4_2':'Executed product category quantity query','sql_Q4_3':'Executed NSW supplier quantity query','model1_process':'Logistic regression overall native process','model1_cv':'Logistic regression cross-validation subprocess','model2_process':'Second model overall native process','model2_cv':'Second model cross-validation subprocess','model1_result':'Native logistic regression fitted on the training partition','model2_result':'Native second model result','model1_performance':'Fixed holdout logistic regression accuracy and confusion matrix','model2_performance':'Fixed holdout second model accuracy and confusion matrix'}

def build(s,final):
 base=ROOT/'submission'/f'Student_{s}'
 missing=[str(base/'evidence'/rel.format(s=s)) for rels in IMAGE_SLOTS.values() for rel in rels if not (base/'evidence'/rel.format(s=s)).exists()]
 if final and missing:raise RuntimeError('Genuine native captures missing: '+', '.join(missing))
 text={}
 for name in ['model','warehouse','query']:text.update(json.loads((base/'report'/f'{name}_sections_draft.json').read_text())['sections'])
 doc=Document(REFERENCE);original=[deepcopy(p._p) for p in doc.paragraphs];body=doc._element.body;section=deepcopy(body.find(q('sectPr')))
 for child in list(body):body.remove(child)
 body.append(section)
 p=doc.add_paragraph('OzMart business analytics report',style='Title');p=doc.add_paragraph(f'STUDENT_{s}_ID',style='Normal')
 headings=[52,57,60,74,84,87,100,115,118,121,123,136,154,156,168,196,198,210,257,262,266,270,282,289,294,298,301,307,310,312,317,321]
 answers={55:'executive_summary',74:'dimensions',115:'implementation',273:'process',297:'predictors',309:'comparison',315:'recommendations',320:'ethics'}
 boxes={125:'Q4_1',157:'Q4_2',199:'Q4_3'}
 for i in range(52,324):
  if i in headings:
   node=deepcopy(original[i])
   pr=node.find('w:pPr',NS)
   if pr is None:pr=E.Element(q('pPr'));node.insert(0,pr)
   if pr.find(q('keepNext')) is None:E.SubElement(pr,q('keepNext'))
   if i in [121,154,196,257,321]:
    pr=node.find('w:pPr',NS)
    if pr is None:pr=E.Element(q('pPr'));node.insert(0,pr)
    if pr.find(q('pageBreakBefore')) is None:E.SubElement(pr,q('pageBreakBefore'))
   body.insert(len(body)-1,node)
  if i in answers:prose(doc,text[answers[i]])
  if i in boxes:sqlbox(doc,original[125],(base/'sql'/f'{boxes[i]}.sql').read_text(),i)
  if i in IMAGE_SLOTS:
   for rel in IMAGE_SLOTS[i]:
    f=base/'evidence'/rel.format(s=s)
    if f.exists():
     p=doc.add_paragraph();p.paragraph_format.keep_with_next=True;p.add_run().add_picture(str(f),width=Inches(6.9));caption=LABELS.get(f.stem,'Dimensional model with actual warehouse columns and keys')
     if f.stem=='model2_result':caption=('Native decision tree fitted on the training partition: one false leaf' if s=='A' else 'Native random forest: first constituent tree shown, with one false leaf')
     prose(doc,caption)
   if i in [138,170,212]:prose(doc,text[{138:'query_4_1',170:'query_4_2',212:'query_4_3'}[i]])
  if i==323:
   refs=json.loads((ROOT/'submission/shared/references/verified_reference_register.json').read_text())['references'];refs=[r.copy() for r in refs if r.get('student',s)==s]
   assert len(refs)==10 and sum(r['task']==1 for r in refs)==5
   if s=='B':
    for r in refs:
     if r['key']=='row_number':r['apa']=r['apa'].replace('(n.d.)','(n.d.-a)')
     if r['key']=='ssis_properties':r['apa']=r['apa'].replace('(n.d.)','(n.d.-b)')
   titles={'kimball_ross':'The data warehouse toolkit: The definitive guide to dimensional modeling','four_steps':'Four-step dimensional design process','grain':'Grain','ssis_properties':'Transformation custom properties','dense_rank':'DENSE_RANK (Transact-SQL)','row_number':'ROW_NUMBER (Transact-SQL)','logistic_operator':'Logistic regression (SVM)','kohavi':'Proceedings of the 14th International Joint Conference on Artificial Intelligence','mehrabi':'A survey on bias and fairness in machine learning','quinlan':'Machine Learning, 1','breiman':'Machine Learning, 45','saito_rehmsmeier':'PLOS ONE, 10'}
   def ref_order(r):
    if r['key']=='kimball_ross':return 'kimball 0'
    if r['key'] in ['four_steps','grain']:return 'kimball 1 '+r['apa']
    if r['apa'].startswith('Microsoft.'):return 'microsoft '+('0' if '(n.d.' in r['apa'] else '1')+r['apa']
    return r['apa'].lower()
   for r in sorted(refs,key=ref_order):
    apa=r['apa'];apa+=(' '+r['url']) if 'https://' not in apa and r['key']!='kimball_ross' else ''
    p=doc.add_paragraph(style='Normal');title=titles[r['key']];before,sep,after=apa.partition(title);assert sep,r['key'];p.add_run(before);p.add_run(title).italic=True;p.add_run(after)
    p.paragraph_format.left_indent=Inches(.5);p.paragraph_format.first_line_indent=Inches(-.5);p.paragraph_format.keep_together=True
   prose(doc,'Generative AI disclosure')
   prose(doc,'Codex provided substantial assistance with design, code, native execution, validation, analysis and report drafting. This use extends beyond proofreading. The accompanying AI disclosure contains visible prompts and AI outputs, code and native execution evidence. Student review and confirmation of course AI and collaboration requirements remain required. No execution outcome is presented as a manual student action.')
 doc.core_properties.author='';doc.core_properties.last_modified_by='';doc.core_properties.title='OzMart business analytics report'
 temp=BytesIO();doc.save(temp)
 # Preserve all unchanged template package parts byte for byte.
 editable={'word/document.xml','word/_rels/document.xml.rels','[Content_Types].xml','docProps/core.xml'}
 old=zipfile.ZipFile(REFERENCE);new=zipfile.ZipFile(temp);outdir=base/'report' if final else ROOT/'work/report_review'/f'Student_{s}';outdir.mkdir(parents=True,exist_ok=True);target=outdir/f'STUDENT_{s}_ID.docx'
 with zipfile.ZipFile(target,'w',compression=zipfile.ZIP_DEFLATED) as z:
  for name in new.namelist():z.writestr(name,old.read(name) if name in old.namelist() and name not in editable else new.read(name))
 result=zipfile.ZipFile(target)
 assert all(result.read(n)==old.read(n) for n in old.namelist() if n not in editable)
 original_sha=hashlib.sha256(REFERENCE.read_bytes()).hexdigest();contract=(ROOT/'docs/REPORT_TEMPLATE_CONTRACT.md').read_text();assert original_sha in contract
 output={'mode':'FINAL' if final else 'REVIEW_INCOMPLETE','student':s,'docx':str(target),'missing_genuine_images':missing,'original_sha256':original_sha,'preserve_only_parts':'PASS','word_counts':{k:len(v.split()) for k,v in text.items()}}
 (outdir/'authoring_validation.json').write_text(json.dumps(output,indent=2)+'\n');print(s,target,'missing native images',len(missing))
if __name__=='__main__':
 a=argparse.ArgumentParser();a.add_argument('--final',action='store_true');args=a.parse_args()
 for s in 'AB':build(s,args.final)
